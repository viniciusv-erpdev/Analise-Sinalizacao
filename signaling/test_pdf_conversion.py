"""Unit tests with a simulated converter; supervisor tests use real OS processes."""
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from io import BytesIO
import importlib.util
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
from threading import Barrier
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from signaling import pdf_conversion as pdf
from signaling.pdf_process import ProcessError, ProcessResult, run_process, runtime_executable

HAS_PYPDF = importlib.util.find_spec("pypdf") is not None


def docx_bytes():
    stream = BytesIO()
    document = Document()
    document.add_heading("Relatório Técnico de Segurança Viária", 0)
    document.add_paragraph("Análise técnica e considerações finais — ç, ã, é.")
    document.save(stream)
    return stream.getvalue()


def fake_pdf():
    from pypdf import PdfWriter
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(stream)
    return stream.getvalue()


@unittest.skipUnless(HAS_PYPDF, "Optional requirements-pdf-poc.txt not installed")
class PDFConversionUnitTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.options = pdf.ConversionOptions(temp_root=self.root)
        self.original = docx_bytes()
        executable = patch.object(pdf, "find_libreoffice", return_value=str(Path(sys.executable).resolve()))
        executable.start()
        self.addCleanup(executable.stop)

    def convert(self):
        return pdf.convert_docx_to_pdf(self.original, options=self.options)

    def emit_pdf(self, command, **kwargs):
        out = Path(command[command.index("--outdir")+1])
        (out / "input.pdf").write_bytes(fake_pdf())
        return ProcessResult(0, b"converted", b"")

    def test_success_has_real_pdf_structure_and_cleans_private_files(self):
        with patch.dict(os.environ, {"POSTGRES_PASSWORD":"SYNTHETIC_SECRET", "DJANGO_SECRET_KEY":"SYNTHETIC_SECRET"}):
            with patch.object(pdf,"run_process",side_effect=self.emit_pdf) as runner:
                converted = self.convert()
        self.assertNotIn("POSTGRES_PASSWORD", runner.call_args.kwargs["env"])
        self.assertNotIn("DJANGO_SECRET_KEY", runner.call_args.kwargs["env"])
        pdf.validate_pdf(converted)
        command = runner.call_args.args[0]
        self.assertIn("--headless", command)
        self.assertIn("pdf:writer_pdf_Export", command)
        self.assertTrue(any(arg.startswith("-env:UserInstallation=file:") for arg in command))
        self.assertEqual(list(self.root.iterdir()), [])
        with ZipFile(BytesIO(self.original)) as original:
            self.assertIsNone(original.testzip())

    def test_actual_project_docx_with_photos_and_chart_is_accepted(self):
        from scripts.pdf_poc import build_sample_docx
        self.original = build_sample_docx()
        with patch.object(pdf,"run_process",side_effect=self.emit_pdf):
            pdf.validate_pdf(self.convert())
        self.assertEqual(list(self.root.iterdir()), [])

    def test_original_content_is_preserved(self):
        before = bytes(self.original)
        with patch.object(pdf,"run_process",side_effect=self.emit_pdf):
            self.convert()
        self.assertEqual(self.original, before)
        self.assertIn(b"PK", before[:2])

    def test_executable_unavailable_is_controlled(self):
        with patch.object(pdf,"find_libreoffice",side_effect=pdf.PDFConversionError("converter_unavailable")):
            with self.assertRaises(pdf.PDFConversionError) as caught:
                self.convert()
        self.assertEqual(caught.exception.code,"converter_unavailable")
        self.assertEqual(list(self.root.iterdir()), [])

    def test_invalid_docx_never_starts_a_converter(self):
        for invalid in (b"", b"not a docx", "not bytes"):
            with self.subTest(value_type=type(invalid).__name__), patch.object(pdf,"run_process") as runner:
                with self.assertRaises(pdf.PDFConversionError):
                    pdf.convert_docx_to_pdf(invalid, options=self.options)
                runner.assert_not_called()

    def test_expansion_and_input_size_are_bounded(self):
        for options in (replace(self.options,max_docx_bytes=1), replace(self.options,max_uncompressed_bytes=10),
                        replace(self.options,max_xml_bytes=10)):
            with self.subTest(options=options), patch.object(pdf,"run_process") as runner:
                with self.assertRaises(pdf.PDFConversionError):
                    pdf.convert_docx_to_pdf(self.original,options=options)
                runner.assert_not_called()

    def test_external_relationships_and_macros_are_rejected(self):
        for name,content in (("word/_rels/extra.rels", b'<Relationships><Relationship TargetMode="External" Target="file:///private"/></Relationships>'),
                             ("word/vbaProject.bin",b"macro")):
            stream = BytesIO(self.original)
            with ZipFile(stream,"a") as archive:
                archive.writestr(name,content)
            with self.subTest(name=name), self.assertRaises(pdf.PDFConversionError):
                pdf.convert_docx_to_pdf(stream.getvalue(),options=self.options)

    def test_xml_entities_and_external_fields_never_reach_converter(self):
        for content in (b'<!DOCTYPE R [<!ENTITY x "expansion">]><R>&x;</R>',
                        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:instrText>INCLUDETEXT file:///private</w:instrText></w:document>'):
            stream = BytesIO(self.original)
            with ZipFile(stream,"a") as archive:
                archive.writestr("word/extra.xml",content)
            with patch.object(pdf,"run_process") as runner:
                with self.assertRaises(pdf.PDFConversionError):
                    pdf.convert_docx_to_pdf(stream.getvalue(),options=self.options)
                runner.assert_not_called()

    def test_timeout_cleans_files(self):
        with patch.object(pdf,"run_process",side_effect=ProcessError("conversion_timeout")):
            with self.assertRaises(pdf.PDFConversionError) as caught:
                self.convert()
        self.assertEqual(caught.exception.code,"conversion_timeout")
        self.assertEqual(list(self.root.iterdir()), [])

    def test_nonzero_exit_and_private_diagnostics(self):
        result = ProcessResult(7,b"PRIVATE CONTENT",b"PRIVATE PATH")
        with patch.object(pdf,"run_process",return_value=result):
            with self.assertRaises(pdf.PDFConversionError) as caught:
                self.convert()
        self.assertEqual(caught.exception.returncode,7)
        self.assertEqual(caught.exception.diagnostics.stderr,b"PRIVATE PATH")
        self.assertNotIn("PRIVATE",str(caught.exception))
        self.assertNotIn("PRIVATE",repr(result))
        self.assertEqual(list(self.root.iterdir()), [])

    def test_success_exit_without_pdf_is_not_success(self):
        with patch.object(pdf,"run_process",return_value=ProcessResult(0,b"",b"")):
            with self.assertRaises(pdf.PDFConversionError) as caught:
                self.convert()
        self.assertEqual(caught.exception.code,"pdf_not_produced")
        self.assertEqual(list(self.root.iterdir()), [])

    def test_corrupt_or_oversized_pdf_is_rejected_and_cleaned(self):
        def emit_bad(command,**kwargs):
            output = Path(command[command.index("--outdir")+1])
            (output / "input.pdf").write_bytes(b"%PDF-1.7\nnot a PDF\n%%EOF")
            return ProcessResult(0,b"",b"")
        with patch.object(pdf,"run_process",side_effect=emit_bad):
            with self.assertRaises(pdf.PDFConversionError):
                self.convert()
        with patch.object(pdf,"run_process",side_effect=self.emit_pdf):
            with self.assertRaises(pdf.PDFConversionError):
                pdf.convert_docx_to_pdf(self.original,options=replace(self.options,max_pdf_bytes=10))
        self.assertEqual(list(self.root.iterdir()), [])

    def test_permission_failure_is_controlled(self):
        with patch.object(pdf.tempfile,"TemporaryDirectory",side_effect=PermissionError):
            with self.assertRaises(pdf.PDFConversionError) as caught:
                self.convert()
        self.assertEqual(caught.exception.code,"temporary_storage_error")

    def test_two_simultaneous_conversions_have_unique_profiles_and_files(self):
        barrier, seen = Barrier(2), []
        def convert(command,**kwargs):
            seen.append((kwargs["cwd"],next(arg for arg in command if arg.startswith("-env:"))))
            self.assertEqual((kwargs["cwd"] / "input.docx").read_bytes(),self.original)
            barrier.wait(timeout=10)
            return self.emit_pdf(command,**kwargs)
        with patch.object(pdf,"run_process",side_effect=convert), ThreadPoolExecutor(2) as executor:
            results = [executor.submit(self.convert) for _ in range(2)]
            self.assertTrue(all(future.result(timeout=15).startswith(b"%PDF-") for future in results))
        self.assertEqual(len({item[0] for item in seen}),2)
        self.assertEqual(len({item[1] for item in seen}),2)
        self.assertEqual(list(self.root.iterdir()), [])

    def test_capacity_is_bounded_without_a_waiting_queue(self):
        pdf._SLOTS.acquire()
        pdf._SLOTS.acquire()
        try:
            with self.assertRaises(pdf.PDFConversionError) as caught:
                self.convert()
            self.assertEqual(caught.exception.code,"conversion_busy")
        finally:
            pdf._SLOTS.release()
            pdf._SLOTS.release()


class ProcessSupervisorTests(unittest.TestCase):
    def run_command(self, script, directory, timeout=5, **limits):
        return run_process([runtime_executable(),"-B","-c",script], cwd=Path(directory), env=os.environ.copy(),
            timeout=timeout,memory_bytes=1024*1024*1024,file_bytes=64*1024*1024,
            workspace_bytes=limits.get("workspace_bytes",256*1024*1024))

    def test_real_process_output_is_captured_and_bounded(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self.run_command("import sys; print('x'*100000); sys.stderr.write('diagnostic')",directory)
        self.assertEqual(result.returncode,0)
        self.assertLessEqual(len(result.stdout),16384)
        self.assertEqual(result.stderr,b"diagnostic")

    def test_real_timeout_terminates_child_process_tree(self):
        script = ("import subprocess,sys,time; from pathlib import Path; "
                  "child=subprocess.Popen([sys.executable,'-B','-c','import time;time.sleep(60)']); "
                  "Path('child.pid').write_text(str(child.pid));time.sleep(60)")
        with tempfile.TemporaryDirectory() as directory:
            started = time.monotonic()
            with self.assertRaises(ProcessError) as caught:
                self.run_command(script,directory,timeout=2)
            self.assertLess(time.monotonic()-started,8, "Process tree outlived the timeout")
            self.assertEqual(caught.exception.code,"conversion_timeout")
            pid = int((Path(directory)/"child.pid").read_text())
            if os.name == "nt":
                import ctypes as c
                from ctypes import wintypes as w
                kernel=c.WinDLL("kernel32",use_last_error=True)
                kernel.OpenProcess.argtypes=[w.DWORD,w.BOOL,w.DWORD]
                kernel.OpenProcess.restype=w.HANDLE
                kernel.WaitForSingleObject.argtypes=[w.HANDLE,w.DWORD]
                kernel.WaitForSingleObject.restype=w.DWORD
                kernel.CloseHandle.argtypes=[w.HANDLE]
                handle=kernel.OpenProcess(0x100000,False,pid)
                if handle:
                    try:
                        self.assertEqual(kernel.WaitForSingleObject(handle,2000),0)
                    finally:
                        kernel.CloseHandle(handle)
            else:
                proc=Path(f"/proc/{pid}/stat")
                self.assertTrue(not proc.exists() or proc.read_text().split()[2]=="Z")

    def test_success_terminates_leftover_children_without_waiting_for_them(self):
        script = ("import subprocess,sys; from pathlib import Path; "
                  "child=subprocess.Popen([sys.executable,'-B','-c','import time;time.sleep(30)']); "
                  "Path('child.pid').write_text(str(child.pid))")
        with tempfile.TemporaryDirectory() as directory:
            started=time.monotonic()
            result=self.run_command(script,directory)
            self.assertEqual(result.returncode,0)
            self.assertLess(time.monotonic()-started,8)

    def test_workspace_overflow_terminates_process(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(ProcessError) as caught:
                self.run_command("from pathlib import Path; import time; Path('large').write_bytes(b'x'*10000);time.sleep(10)",
                                 directory,workspace_bytes=100)
            self.assertEqual(caught.exception.code,"temporary_storage_limit")


class DOCXRegressionWithoutDatabaseTests(unittest.TestCase):
    def test_real_generator_preserves_report_structure_styles_images_and_original_fields(self):
        from scripts.pdf_poc import build_sample_docx
        from xml.etree import ElementTree
        with patch("django.db.backends.base.base.BaseDatabaseWrapper.cursor",
                   side_effect=AssertionError("No SQL permitted in PDF PoC")):
            original = build_sample_docx()
        document = Document(BytesIO(original))
        text = "\n".join([p.text for p in document.paragraphs] +
                         [cell.text for table in document.tables for row in table.rows for cell in row.cells])
        for expected in ("Relatório Técnico de Segurança Viária", "Análise técnica e considerações finais",
                         "Autor sintético da PoC", "Data de emissão", "Escola fictícia", "SINTETICO-30",
                         "40 km/h", "Vistoriador sintético", "Problema sintético de travessia"):
            self.assertIn(expected,text)
        self.assertGreaterEqual(len(document.inline_shapes),1)  # actual Matplotlib chart
        self.assertEqual(document.styles["Normal"].font.name,"Calibri")
        self.assertEqual(document.styles["Title"].font.name,"Arial")
        self.assertTrue(document.sections[0].footer.paragraphs[0].text)
        with ZipFile(BytesIO(original)) as archive:
            self.assertGreaterEqual(len([name for name in archive.namelist() if name.startswith("word/media/")]),5)
            for name in archive.namelist():
                if name.endswith(".xml"):
                    ElementTree.fromstring(archive.read(name))
