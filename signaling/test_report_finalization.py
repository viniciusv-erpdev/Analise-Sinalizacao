"""No-database tests for immutable inputs/storage/rendering, not PostgreSQL proof."""
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4

from django.core.exceptions import ImproperlyConfigured, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, override_settings
from docx import Document
from PIL import Image

from signaling.report_snapshots import (prepare_photos, render_snapshot, digest, validate_form, OFFICIAL_TIMEZONE)
from signaling.report_storage import PrivateReportStorage, ArtifactStorageError
from signaling.report_finalization import require_enabled, FinalizationError


def photo_upload():
    stream = BytesIO()
    Image.new("RGB", (20, 30), "red").save(stream, "PNG")
    return SimpleUploadedFile("photo.png", stream.getvalue(), content_type="image/png")


def sample_snapshot():
    return {
        "schema_version": 1, "official_number": "RTSV-2026-0001",
        "issued_at": "2026-12-31T22:30:00-03:00",
        "source": {"point": {"id": 123, "search_radius_meters": 50}},
        "form": {"inspection_address": "Synthetic street", "inspection_date": "2026-12-31",
                 "inspector_name": "Inspector", "pgts": ["School"], "technical_analysis": "Final analysis"},
        "author": {"username": "frozen-user", "full_name": "Frozen Author"},
        "survey": {"summary": {}, "accidents": [], "interventions": [], "problems": []},
        "photos": [],
    }


class ReportSnapshotTests(SimpleTestCase):
    def test_snapshot_renders_without_database_or_session_and_preserves_input(self):
        snapshot = sample_snapshot()
        before = deepcopy(snapshot)
        document = Document(BytesIO(render_snapshot(snapshot, [])))
        text = "\n".join(p.text for p in document.paragraphs)
        table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("RTSV-2026-0001", text)
        self.assertIn("Frozen Author (frozen-user)", text)
        self.assertIn("Final analysis", text)
        self.assertIn("31/12/2026", text)
        self.assertIn("31/12/2026", table_text)
        self.assertEqual(snapshot, before)

    def test_official_year_is_sao_paulo_even_when_utc_is_next_year(self):
        instant = datetime(2027, 1, 1, 1, 30, tzinfo=timezone.utc)
        self.assertEqual(instant.astimezone(OFFICIAL_TIMEZONE).year, 2026)
        self.assertEqual(instant.astimezone(OFFICIAL_TIMEZONE).isoformat(), "2026-12-31T22:30:00-03:00")

    def test_unknown_schema_rejected(self):
        snapshot = sample_snapshot()
        snapshot["schema_version"] = 999
        with self.assertRaises(ValueError):
            render_snapshot(snapshot, [])

    def test_photo_original_order_hash_dimensions_and_bytes(self):
        uploads = [photo_upload(), photo_upload()]
        originals = [photo.read() for photo in uploads]
        result = prepare_photos(uploads)
        self.assertEqual([meta["ordinal"] for meta, _ in result], [0, 1])
        self.assertEqual([data for _, data in result], originals)
        self.assertEqual(result[0][0]["width"], 20)
        self.assertEqual(result[0][0]["height"], 30)
        self.assertEqual(result[0][0]["content_type"], "image/png")
        import hashlib
        self.assertEqual(result[0][0]["sha256"], hashlib.sha256(originals[0]).hexdigest())

    def test_invalid_photo_and_quantity_and_size_rejected(self):
        invalid = SimpleUploadedFile("x.png", b"not an image", content_type="image/png")
        with self.assertRaises(ValidationError):
            prepare_photos([invalid])
        with self.assertRaises(ValidationError):
            prepare_photos([photo_upload()] * 11)
        with patch("signaling.report_snapshots.MAX_PHOTOS_TOTAL_BYTES", 1):
            with self.assertRaises(ValidationError):
                prepare_photos([photo_upload()])
        with patch("signaling.report_snapshots.MAX_PHOTO_PIXELS", 1):
            with self.assertRaises(ValidationError):
                prepare_photos([photo_upload()])

    def test_form_revalidates_required_fields_dates_and_limits(self):
        fields = sample_snapshot()["form"]
        self.assertEqual(validate_form(fields)["technical_analysis"], "Final analysis")
        for changes in ({"inspection_date": "bad"}, {"inspection_address": ""}, {"technical_analysis": "x" * 10001}):
            with self.assertRaises(ValidationError):
                validate_form(fields | changes)

    def test_canonical_hash_stable_and_rejects_nonfinite(self):
        self.assertEqual(digest({"a": 1, "b": 2}), digest({"b": 2, "a": 1}))
        with self.assertRaises(ValueError):
            digest({"x": float("nan")})

    @override_settings(REPORT_FINALIZATION_ENABLED=False)
    def test_disabled_by_default(self):
        with self.assertRaises(FinalizationError) as caught:
            require_enabled()
        self.assertEqual(caught.exception.code, "finalization_disabled")


class PrivateStorageTests(SimpleTestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.override = override_settings(REPORT_PRIVATE_ROOT=self.temp.name)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.storage = PrivateReportStorage()
        self.key = f"{uuid4()}/{uuid4()}/{uuid4()}.pdf"

    def test_private_immutable_file_and_integrity(self):
        import hashlib
        data = b"synthetic artifact"
        self.storage.save(self.key, data)
        artifact = SimpleNamespace(storage_key=self.key, size=len(data), sha256=hashlib.sha256(data).hexdigest())
        self.assertEqual(self.storage.read(artifact), data)
        self.assertEqual(list(self.storage.keys()), [self.key])
        with self.assertRaises(ArtifactStorageError):
            self.storage.save(self.key, b"overwrite")
        self.assertEqual(self.storage.read(artifact), data)
        Path(self.temp.name, self.key).write_bytes(b"tampered")
        with self.assertRaises(ArtifactStorageError):
            self.storage.read(artifact)
        with self.assertRaises(NotImplementedError):
            self.storage.url(self.key)

    def test_paths_and_repository_root_rejected(self):
        for key in ("../evil.pdf", "C:/evil.pdf", "photo.png"):
            with self.assertRaises(ArtifactStorageError):
                self.storage.save(key, b"x")
        from django.conf import settings
        with override_settings(REPORT_PRIVATE_ROOT=str(settings.BASE_DIR / "private")):
            with self.assertRaises(ImproperlyConfigured):
                PrivateReportStorage()


class FinalizationErrorStatusTests(SimpleTestCase):
    def test_authentication_and_permission_errors_remain_distinct_from_reservation_conflict(self):
        from signaling import reservations as r
        from signaling.report_finalization import _error_status
        self.assertEqual(_error_status(r.ReservationAuthenticationRequired()), 401)
        self.assertEqual(_error_status(r.ReservationPermissionDenied()), 403)
        self.assertEqual(_error_status(r.ReservationExpired()), 409)
        self.assertEqual(_error_status(r.ReservationWaypointNotFound()), 404)
        self.assertEqual(_error_status(FinalizationError("conflict", 409)), 409)


class ArtifactRegistrationTests(SimpleTestCase):
    """Real temporary files; DB/commit injected, never a PostgreSQL substitute."""
    def setUp(self):
        PrivateStorageTests.setUp(self)
        from signaling import report_finalization as service
        from contextlib import nullcontext
        self.service = service
        self.attempt = SimpleNamespace(pk=uuid4(), version_id=uuid4(), execution_id=uuid4(),
            version=SimpleNamespace(version=1, report=SimpleNamespace(official_number="RTSV-2026-0001")))
        for target, kwargs in (
            ("transaction.atomic", {"side_effect": lambda: nullcontext()}),
            ("_assert_run", {}),
            ("ReportArtifact.objects.create", {"side_effect": lambda **fields: SimpleNamespace(**fields)}),
        ):
            mocked = patch("signaling.report_finalization." + target, **kwargs)
            result = mocked.start()
            self.addCleanup(mocked.stop)
            if target == "ReportArtifact.objects.create":
                self.create = result

    def store(self):
        return self.service._store(self.attempt, self.storage, "PHOTO", b"original", metadata={"format":"PNG"})

    def test_file_is_verified_before_registration(self):
        def create(**fields):
            artifact = SimpleNamespace(**fields)
            self.assertEqual(self.storage.read(artifact), b"original")
            return artifact
        self.create.side_effect = create
        self.assertEqual(self.storage.read(self.store()), b"original")

    def test_read_failure_after_save_keeps_unregistered_file_and_original_error(self):
        primary = ArtifactStorageError("artifact_missing")
        with patch.object(self.storage, "read", side_effect=primary), patch.object(
                self.storage, "delete", side_effect=RuntimeError("cleanup failure")) as cleanup:
            with self.assertRaises(ArtifactStorageError) as caught:
                self.store()
        self.assertIs(caught.exception, primary)
        self.create.assert_not_called()
        cleanup.assert_not_called()
        self.assertEqual(len(list(self.storage.keys())), 1)

    def test_size_and_hash_corruption_rejected_before_registration(self):
        original_save = self.storage.save
        for corrupt in (b"short", b"modified"):
            def save(key, data):
                original_save(key, data)
                Path(self.temp.name, key).write_bytes(corrupt)
            with patch.object(self.storage, "save", side_effect=save):
                with self.assertRaises(ArtifactStorageError):
                    self.store()
        self.create.assert_not_called()
        self.assertEqual(len(list(self.storage.keys())), 2)

    def test_registration_failure_preserves_orphan(self):
        failure = RuntimeError("registration failed")
        self.create.side_effect = failure
        with self.assertRaises(RuntimeError) as caught:
            self.store()
        self.assertIs(caught.exception, failure)
        self.assertEqual(len(list(self.storage.keys())), 1)

    def test_ambiguous_commit_preserves_registered_file(self):
        from contextlib import contextmanager
        @contextmanager
        def ambiguous():
            yield
            raise RuntimeError("commit acknowledgement lost")
        with patch.object(self.service.transaction, "atomic", side_effect=ambiguous):
            with self.assertRaisesRegex(RuntimeError, "acknowledgement lost"):
                self.store()
        fields = self.create.call_args.kwargs
        self.assertEqual(self.storage.read(SimpleNamespace(**fields)), b"original")
