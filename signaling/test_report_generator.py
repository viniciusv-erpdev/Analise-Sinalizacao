from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile, is_zipfile
from xml.etree import ElementTree

from django.test import TestCase
from docx import Document
from PIL import Image

from signaling.models import SignalingIntervention, SignalingPoint
from signaling.report_generator import (
    INTERVENTION_ICON_DIRECTORY,
    InvalidReportImage,
    REPORT_FOOTER_LINES,
    REPORT_LOGO_PATH,
    build_accident_types_chart,
    generate_signaling_report,
    svg_to_png_buffer,
)


class SignalingReportGeneratorTests(TestCase):
    def setUp(self):
        self.point = SignalingPoint.objects.create(
            latitude=-21.17,
            longitude=-47.81,
            status=SignalingPoint.Status.INCOMPLETE,
            search_radius_meters=100,
        )
        self.survey = {
            "radius_meters": 100,
            "status": {"value": "INCOMPLETE", "label": "Incompleta"},
            "total_accidents": 4,
            "summary": {
                "fatal": 1,
                "non_fatal": 3,
                "by_type": {
                    "Atropelamento": 1,
                    "Choque": 0,
                    "Colisão": 2,
                    "Não disponível": 0,
                    "Outros": 1,
                },
            },
            "interventions": [],
            "accidents": [],
        }

    def generate(self, **kwargs) -> BytesIO:
        return generate_signaling_report(
            self.point,
            self.survey,
            kwargs.get("form_data"),
            kwargs.get("photos", ()),
        )

    def document_text(self, report: BytesIO) -> str:
        document = Document(report)
        paragraph_text = [paragraph.text for paragraph in document.paragraphs]
        table_text = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        return "\n".join(paragraph_text + table_text)

    def find_table(self, document, headers):
        expected = list(headers)
        for table in document.tables:
            if table.rows and [cell.text for cell in table.rows[0].cells] == expected:
                return table
        self.fail(f"Tabela com cabeçalho {expected!r} não encontrada.")

    def test_generates_valid_docx_in_bytes_io(self):
        report = self.generate()

        self.assertIsInstance(report, BytesIO)
        self.assertTrue(is_zipfile(report))
        self.assertIn("Relatório do local", self.document_text(report))

    def test_all_docx_xml_parts_are_well_formed_and_word_compatible(self):
        self.survey["accidents"] = [{
            "id": "A-10",
            "date": "10/09/2026",
            "accident_type": "Colisão",
            "record_type": "SINISTRO NAO FATAL",
            "street": "AVENIDA TESTE",
            "distance_meters": 42.0,
            "modes": [{"name": "Automóvel", "quantity": 1}],
        }]
        self.survey["interventions"] = [{
            "type": SignalingIntervention.Type.TRAFFIC_LIGHT,
            "type_label": "Semáforo",
            "condition_label": "Adequada",
            "notes": "Boa visibilidade.",
        }]
        photos = []
        for color in ("red", "green", "blue"):
            photo = BytesIO()
            Image.new("RGB", (100, 60), color).save(photo, format="PNG")
            photo.seek(0)
            photos.append(photo)

        generated_bytes = self.generate(photos=photos).getvalue()

        with ZipFile(BytesIO(generated_bytes)) as archive:
            self.assertIn("[Content_Types].xml", archive.namelist())
            self.assertIn("word/document.xml", archive.namelist())
            for name in archive.namelist():
                if name.endswith(".xml"):
                    ElementTree.fromstring(archive.read(name))

            document_xml = archive.read("word/document.xml")
            self.assertNotIn(b"<w:start", document_xml)
            self.assertNotIn(b"<w:end", document_xml)
            first_cell_properties = document_xml.split(b"<w:tcPr>", 1)[1].split(
                b"</w:tcPr>", 1
            )[0]
            self.assertLess(
                first_cell_properties.index(b"<w:tcBorders>"),
                first_cell_properties.index(b"<w:shd"),
            )

        reopened = Document(BytesIO(generated_bytes))
        self.assertIn("Relatório do local", reopened.paragraphs[0].text)

    def test_includes_form_and_survey_data(self):
        report = self.generate(form_data={
            "inspection_address": "Avenida Teste, 100",
            "occurrence_date": date(2026, 8, 1),
            "inspection_date": date(2026, 8, 2),
            "inspector_name": "Responsável Teste",
        })
        text = self.document_text(report)

        self.assertIn("Avenida Teste, 100", text)
        self.assertIn("01/08/2026", text)
        self.assertIn("Responsável Teste", text)
        self.assertIn("100 m", text)
        self.assertIn("Incompleta", text)
        self.assertIn("4", text)

    def test_includes_the_same_validated_filter_summary(self):
        self.survey["filters"] = {
            "period": "2025 — Janeiro",
            "categories": "Atropelamento, Colisão",
            "gravity": "Somente fatais",
        }

        text = self.document_text(self.generate())

        self.assertIn("Filtros aplicados", text)
        self.assertIn("2025 — Janeiro", text)
        self.assertIn("Atropelamento, Colisão", text)
        self.assertIn("Somente fatais", text)
        self.assertIn("100 m", text)

    def test_includes_intervention_and_notes(self):
        self.survey["interventions"] = [{
            "id": 1,
            "type": SignalingIntervention.Type.TRAFFIC_LIGHT,
            "type_label": "Semáforo",
            "condition": SignalingIntervention.Condition.OK,
            "condition_label": "Adequada",
            "notes": "Lado norte.",
        }]

        report = self.generate()
        text = self.document_text(report)

        self.assertIn("Semáforo", text)
        self.assertIn("Adequada", text)
        self.assertIn("Lado norte.", text)

    def test_includes_study_reason_and_objective(self):
        report = self.generate(form_data={
            "study_reason": "Reduzir conflitos no cruzamento.",
            "study_objective": "Avaliar a sinalização existente.",
        })
        text = self.document_text(report)

        self.assertIn("Motivo do estudo", text)
        self.assertIn("Reduzir conflitos no cruzamento.", text)
        self.assertIn("Objetivo", text)
        self.assertIn("Avaliar a sinalização existente.", text)

    def test_includes_related_accidents_table(self):
        self.survey["accidents"] = [{
            "id": "A-10",
            "date": "10/09/2026",
            "accident_type": "Colisão",
            "record_type": "SINISTRO NAO FATAL",
            "street": "AVENIDA TESTE",
            "distance_meters": 42.04,
            "modes": [
                {"name": "Automóvel", "quantity": 1},
                {"name": "Motocicleta", "quantity": 2},
            ],
        }]

        report = self.generate()
        text = self.document_text(report)

        for header in ("ID", "DATA", "TIPO", "GRAVIDADE", "LOCAL", "DISTÂNCIA", "MODAIS"):
            self.assertIn(header, text)
        self.assertIn("A-10", text)
        self.assertIn("10/09/2026", text)
        self.assertIn("AVENIDA TESTE", text)
        self.assertIn("42.0 m", text)
        self.assertIn("Motocicleta: 2", text)

        document = Document(report)
        self.find_table(
            document,
            ("ID", "DATA", "TIPO", "GRAVIDADE", "LOCAL", "DISTÂNCIA", "MODAIS"),
        )

    def test_includes_institutional_header_footer_and_global_fonts(self):
        document = Document(self.generate())
        header = document.sections[0].header
        footer = document.sections[0].footer

        embedded_images = [
            part.blob
            for part in header.part.related_parts.values()
            if hasattr(part, "blob")
        ]
        self.assertIn(REPORT_LOGO_PATH.read_bytes(), embedded_images)
        self.assertEqual(
            "\n".join(paragraph.text for paragraph in footer.paragraphs),
            "\n".join(REPORT_FOOTER_LINES),
        )
        self.assertEqual(document.styles["Normal"].font.name, "Calibri")
        self.assertEqual(document.styles["Title"].font.name, "Arial")

    def test_uses_styled_information_and_intervention_tables(self):
        self.survey["interventions"] = [{
            "type": SignalingIntervention.Type.TRAFFIC_LIGHT,
            "type_label": "Semáforo",
            "condition_label": "Adequada",
            "notes": "Boa visibilidade.",
        }]
        document = Document(self.generate())

        information_table = next(
            table for table in document.tables
            if table.cell(0, 0).text == "Endereço da vistoria"
        )
        intervention_table = self.find_table(
            document,
            ("ÍCONE", "INTERVENÇÃO", "STATUS", "OBSERVAÇÃO"),
        )

        self.assertEqual(len(information_table.rows), 6)
        self.assertEqual(intervention_table.cell(1, 1).text, "Semáforo")
        self.assertEqual(intervention_table.cell(1, 2).text, "Adequada")
        self.assertEqual(intervention_table.cell(1, 3).text, "Boa visibilidade.")

    def test_chart_is_generated_as_png(self):
        chart = build_accident_types_chart(self.survey)

        self.assertTrue(chart.getvalue().startswith(b"\x89PNG\r\n\x1a\n"))

    def test_converts_existing_svg_to_png(self):
        icon = svg_to_png_buffer(INTERVENTION_ICON_DIRECTORY / "traffic-light.svg")

        self.assertTrue(icon.getvalue().startswith(b"\x89PNG\r\n\x1a\n"))

    def test_embeds_chart_and_intervention_icon(self):
        self.survey["interventions"] = [{
            "type": SignalingIntervention.Type.TRAFFIC_LIGHT,
            "type_label": "Semáforo",
            "condition_label": "Adequada",
            "notes": "",
        }]
        report = self.generate()

        with ZipFile(report) as archive:
            media_files = [
                name
                for name in archive.namelist()
                if name.startswith("word/media/")
            ]

        self.assertGreaterEqual(len(media_files), 2)

    @patch("signaling.report_generator.svg_to_png_buffer", side_effect=ValueError)
    def test_icon_conversion_failure_does_not_abort_report(self, _convert_mock):
        self.survey["interventions"] = [{
            "type": SignalingIntervention.Type.TRAFFIC_LIGHT,
            "type_label": "Semáforo",
            "condition_label": "Adequada",
            "notes": "Visível",
        }]

        with self.assertLogs("signaling.report_generator", level="WARNING"):
            report = self.generate()
        text = self.document_text(report)

        self.assertTrue(is_zipfile(report))
        self.assertIn("Semáforo", text)
        self.assertIn("Adequada", text)
        self.assertIn("Visível", text)

    def test_embeds_temporary_png_photo(self):
        photo = BytesIO()
        Image.new("RGB", (100, 60), "blue").save(photo, format="PNG")
        photo.seek(0)

        report = self.generate(photos=[photo])

        with ZipFile(report) as archive:
            media_files = [
                name
                for name in archive.namelist()
                if name.startswith("word/media/")
            ]
        self.assertGreaterEqual(len(media_files), 2)
        self.assertFalse(Path("chart.png").exists())

    def test_places_an_odd_number_of_photos_in_two_column_grid(self):
        photos = []
        for color in ("red", "green", "blue"):
            photo = BytesIO()
            Image.new("RGB", (100, 60), color).save(photo, format="PNG")
            photo.seek(0)
            photos.append(photo)

        document = Document(self.generate(photos=photos))
        photo_table = document.tables[-1]

        self.assertEqual(len(photo_table.rows), 2)
        self.assertEqual(len(photo_table.columns), 2)
        self.assertEqual(len(photo_table._element.xpath(".//a:blip")), 3)

    def test_rejects_invalid_photo_with_friendly_error(self):
        with self.assertRaisesMessage(InvalidReportImage, "JPEG ou PNG válida"):
            self.generate(photos=[BytesIO(b"not an image")])
