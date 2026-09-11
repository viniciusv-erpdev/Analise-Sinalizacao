"""Gera relatórios DOCX transitórios para pontos de sinalização."""

from collections.abc import Iterable, Mapping
import base64
from datetime import date, datetime
from io import BytesIO
import logging
from pathlib import Path
from typing import BinaryIO

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from django.conf import settings
from django.utils import timezone
from PIL import Image, UnidentifiedImageError
from resvg_py import svg_to_bytes

from signaling.models import SignalingPoint
from signaling.intervention_icons import get_intervention_icon_filename
from signaling.docx_styles import (
    NAVY,
    SOFT_GRAY,
    WHITE,
    apply_run_font,
    configure_document_styles,
    format_table_body,
    format_table_header,
    set_cell_background,
    set_cell_border,
    set_cell_margins,
    style_cell_text,
)


logger = logging.getLogger(__name__)

INTERVENTION_ICON_DIRECTORY = (
    Path(__file__).resolve().parent.parent / "static" / "icons" / "signaling"
)
REPORT_LOGO_PATH = Path(settings.BASE_DIR) / "static" / "logo-tipo" / "rpmobi_logo.jpg"
REPORT_FOOTER_LINES = (
    "Rua General Câmara, 2910 – Vila Recreio – PABX (16) 3934-9500",
    "14060-582 – Ribeirão Preto/SP - www.ribeiraopreto.sp.gov.br/rpmobi",
)

REPORT_CATEGORY_LABELS = (
    "Atropelamento",
    "Choque",
    "Colisão",
    "Não disponível",
    "Outros",
)


class InvalidReportImage(ValueError):
    """Indica que uma foto não pode ser inserida no relatório."""


def svg_to_png_buffer(svg_path: Path) -> BytesIO:
    """Converte um SVG existente em PNG, inteiramente em memória."""
    png_buffer = BytesIO(svg_to_bytes(svg_path=str(svg_path)))
    png_buffer.seek(0)
    return png_buffer


def build_accident_types_chart(survey: Mapping[str, object]) -> BytesIO:
    """Cria em memória o gráfico de tipos de sinistro do levantamento."""
    summary = survey.get("summary", {})
    counts_by_type = summary.get("by_type", {}) if isinstance(summary, Mapping) else {}
    quantities = [int(counts_by_type.get(label, 0)) for label in REPORT_CATEGORY_LABELS]

    figure, axis = plt.subplots(figsize=(7.0, 3.4))
    try:
        bars = axis.bar(
            REPORT_CATEGORY_LABELS,
            quantities,
            color=("#dc3545", "#198754", "#0d6efd", "#ffc107", "#0dcaf0"),
            width=0.62,
        )
        axis.set_ylabel("Quantidade")
        axis.tick_params(axis="x", labelrotation=16, length=0)
        axis.tick_params(axis="y", colors="#6c757d", length=0)
        axis.spines[["top", "right"]].set_visible(False)
        axis.spines[["left", "bottom"]].set_color("#d9d9d9")
        axis.set_axisbelow(True)
        axis.grid(axis="y", color="#e9ecef", linewidth=0.7)
        axis.bar_label(bars, padding=3, fontsize=8, color="#003366")
        axis.margins(y=0.14)
        figure.tight_layout()

        chart_buffer = BytesIO()
        figure.savefig(
            chart_buffer,
            format="png",
            dpi=300,
            bbox_inches="tight",
            transparent=False,
        )
        chart_buffer.seek(0)
        return chart_buffer
    finally:
        plt.close(figure)


def build_accident_types_chart_data_uri(survey: Mapping[str, object]) -> str:
    """Transforma o mesmo gráfico PNG usado no Word em uma data URI HTML."""
    chart_buffer = build_accident_types_chart(survey)
    encoded_chart = base64.b64encode(chart_buffer.getvalue()).decode("ascii")
    return f"data:image/png;base64,{encoded_chart}"


def generate_signaling_report(
    signaling_point: SignalingPoint,
    survey: Mapping[str, object],
    form_data: Mapping[str, object] | None = None,
    photos: Iterable[BinaryIO] = (),
) -> BytesIO:
    """Retorna um DOCX em memória sem persistir formulário, fotos ou arquivo."""
    temporary_data = form_data or {}
    document = Document()
    configure_document_styles(document)
    _add_header(document)
    _add_footer(document)
    document.add_heading("Relatório do local", level=0)

    document.add_heading("Informações do estudo", level=1)
    _add_information_table(document, (
        ("Endereço da vistoria", _display_value(temporary_data.get("inspection_address"))),
        ("Data de geração do relatório", timezone.localdate().strftime("%d/%m/%Y")),
        ("Data de ocorrência", _display_value(temporary_data.get("occurrence_date"))),
        ("Data da vistoria", _display_value(temporary_data.get("inspection_date"))),
        ("Motivo do estudo", _display_value(temporary_data.get("study_reason"))),
        ("Responsável pela vistoria", _display_value(temporary_data.get("inspector_name"))),
    ))

    document.add_heading("Objetivo", level=1)
    _add_objective_block(
        document,
        _display_value(temporary_data.get("study_objective")),
    )

    summary = survey.get("summary", {})
    if not isinstance(summary, Mapping):
        summary = {}
    status = survey.get("status", {})
    status_label = status.get("label") if isinstance(status, Mapping) else None

    document.add_heading("Resumo", level=1)
    _add_summary_cards(
        document,
        total=str(survey.get("total_accidents", 0)),
        fatal=str(summary.get("fatal", 0)),
        non_fatal=str(summary.get("non_fatal", 0)),
        radius=f"{signaling_point.search_radius_meters} m",
        status=_display_value(status_label),
    )

    document.add_heading("Tipos de sinistros", level=1)
    chart_buffer = build_accident_types_chart(survey)
    document.add_picture(chart_buffer, width=Cm(16))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("Sinistros relacionados", level=1)
    accidents = survey.get("accidents", [])
    if accidents:
        _add_related_accidents_table(document, accidents)
    else:
        document.add_paragraph(
            "Nenhum sinistro encontrado dentro do raio analisado."
        )

    document.add_heading("Intervenções", level=1)
    interventions = survey.get("interventions", [])
    if interventions:
        _add_interventions_table(document, interventions)
    else:
        document.add_paragraph("Nenhuma intervenção cadastrada.")

    document.add_heading("Fotos do local", level=1)
    photo_buffers = [_validated_photo_buffer(photo) for photo in photos]
    if photo_buffers:
        _add_photos_grid(document, photo_buffers)
    else:
        document.add_paragraph("Nenhuma foto informada.")

    report_buffer = BytesIO()
    document.save(report_buffer)
    report_buffer.seek(0)
    return report_buffer


def _add_header(document: DocumentObject) -> None:
    header = document.sections[0].header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run().add_picture(str(REPORT_LOGO_PATH), width=Cm(4.0))


def _add_footer(document: DocumentObject) -> None:
    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Cm(0)
    paragraph.paragraph_format.space_after = Cm(0)
    for index, line in enumerate(REPORT_FOOTER_LINES):
        if index:
            paragraph.add_run().add_break()
        apply_run_font(paragraph.add_run(line), name="Verdana", size=8)


def _add_information_table(
    document: DocumentObject,
    rows: Iterable[tuple[str, str]],
) -> None:
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_background(cells[0], SOFT_GRAY)
        set_cell_background(cells[1], WHITE)
        for cell in cells:
            set_cell_border(cell)
            set_cell_margins(cell, top=110, bottom=110)
        style_cell_text(cells[0], name="Arial", color=NAVY, bold=True)
        style_cell_text(cells[1])


def _add_objective_block(document: DocumentObject, objective: str) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = objective
    set_cell_background(cell, SOFT_GRAY)
    set_cell_border(cell)
    set_cell_margins(cell, top=140, start=140, bottom=140, end=140)
    style_cell_text(cell)


def _add_summary_cards(
    document: DocumentObject,
    *,
    total: str,
    fatal: str,
    non_fatal: str,
    radius: str,
    status: str,
) -> None:
    table = document.add_table(rows=3, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ("TOTAL DE SINISTROS", "FATAIS", "NÃO FATAIS", "RAIO")
    values = (total, fatal, non_fatal, radius)
    for cell, label in zip(table.rows[0].cells, labels):
        cell.text = label
        set_cell_background(cell, SOFT_GRAY)
        set_cell_border(cell)
        set_cell_margins(cell)
        style_cell_text(
            cell,
            name="Arial",
            size=9,
            color=NAVY,
            bold=True,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for cell, value in zip(table.rows[1].cells, values):
        cell.text = value
        set_cell_background(cell, WHITE)
        set_cell_border(cell)
        set_cell_margins(cell, top=130, bottom=130)
        style_cell_text(
            cell,
            name="Arial",
            size=14,
            color=NAVY,
            bold=True,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )

    status_cell = table.rows[2].cells[0]
    for other_cell in table.rows[2].cells[1:]:
        status_cell = status_cell.merge(other_cell)
    status_cell.text = f"Status da sinalização: {status}"
    set_cell_background(status_cell, SOFT_GRAY)
    set_cell_border(status_cell)
    set_cell_margins(status_cell)
    style_cell_text(
        status_cell,
        name="Arial",
        color=NAVY,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )


def _add_interventions_table(
    document: DocumentObject,
    interventions: Iterable[Mapping[str, object]],
) -> None:
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    column_widths = (1.5, 4.0, 2.5, 8.0)
    headers = ("ÍCONE", "INTERVENÇÃO", "STATUS", "OBSERVAÇÃO")
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for cell, width in zip(table.rows[0].cells, column_widths):
        cell.width = Cm(width)
    format_table_header(table, left_aligned_columns=(1, 3))

    for intervention in interventions:
        icon_cell, type_cell, status_cell, notes_cell = table.add_row().cells
        for cell, width in zip(
            (icon_cell, type_cell, status_cell, notes_cell),
            column_widths,
        ):
            cell.width = Cm(width)
        icon_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        icon_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        icon_name = get_intervention_icon_filename(
            str(intervention.get("type", ""))
        )
        if icon_name:
            try:
                icon_buffer = svg_to_png_buffer(
                    INTERVENTION_ICON_DIRECTORY / icon_name
                )
                icon_cell.paragraphs[0].add_run().add_picture(
                    icon_buffer,
                    width=Cm(1.1),
                )
            except (OSError, ValueError):
                logger.warning(
                    "Não foi possível converter o ícone %s",
                    icon_name,
                    exc_info=True,
                )
        type_cell.text = _display_value(intervention.get("type_label"))
        status_cell.text = _display_value(intervention.get("condition_label"))
        notes_cell.text = _display_value(intervention.get("notes"))

    format_table_body(table, font_size=9)


def _add_related_accidents_table(
    document: DocumentObject,
    accidents: Iterable[Mapping[str, object]],
) -> None:
    headers = ("ID", "DATA", "TIPO", "GRAVIDADE", "LOCAL", "DISTÂNCIA", "MODAIS")
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    column_widths = (1.3, 1.8, 2.0, 2.6, 4.1, 1.6, 3.4)
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for cell, width in zip(table.rows[0].cells, column_widths):
        cell.width = Cm(width)
    format_table_header(table, left_aligned_columns=(4, 6))

    for accident in accidents:
        modes = accident.get("modes", [])
        modes_text = ", ".join(
            f"{mode.get('name')}: {mode.get('quantity')}"
            for mode in modes
        ) or "Não disponível"
        distance = accident.get("distance_meters")
        try:
            distance_text = f"{float(distance):.1f} m"
        except (TypeError, ValueError):
            distance_text = "Não disponível"
        values = (
            accident.get("id"),
            accident.get("date"),
            accident.get("accident_type"),
            accident.get("record_type"),
            accident.get("street"),
            distance_text,
            modes_text,
        )
        cells = table.add_row().cells
        for cell, value, width in zip(cells, values, column_widths):
            cell.text = _display_value(value)
            cell.width = Cm(width)

    format_table_body(table, font_size=8)


def _add_photos_grid(
    document: DocumentObject,
    photo_buffers: list[BytesIO],
) -> None:
    row_count = (len(photo_buffers) + 1) // 2
    table = document.add_table(rows=row_count, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(8)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_background(cell, WHITE)
            set_cell_border(cell)
            set_cell_margins(cell, top=110, start=110, bottom=110, end=110)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for index, photo_buffer in enumerate(photo_buffers):
        row_index, column_index = divmod(index, 2)
        cell = table.cell(row_index, column_index)
        cell.paragraphs[0].add_run().add_picture(photo_buffer, width=Cm(7.2))


def _validated_photo_buffer(photo: BinaryIO) -> BytesIO:
    try:
        photo.seek(0)
        image_bytes = photo.read()
        photo.seek(0)
        image_buffer = BytesIO(image_bytes)
        with Image.open(image_buffer) as image:
            if image.format not in {"JPEG", "PNG"}:
                raise InvalidReportImage("Envie fotos nos formatos JPEG ou PNG.")
            image.verify()
    except (AttributeError, OSError, UnidentifiedImageError) as error:
        raise InvalidReportImage("A foto enviada não é uma imagem JPEG ou PNG válida.") from error

    image_buffer.seek(0)
    return image_buffer


def _display_value(value: object) -> str:
    if value is None or value == "":
        return "Não informado"
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    return str(value)
