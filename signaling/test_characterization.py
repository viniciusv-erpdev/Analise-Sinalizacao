from io import BytesIO
from unittest.mock import patch
from xml.etree import ElementTree
from zipfile import ZipFile, is_zipfile

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from docx import Document

from signaling.forms import PGTListField, SignalingReportForm
from signaling.models import SignalingIntervention, SignalingPoint
from signaling.surveys import ANALYSIS_SESSION_KEY


class CharacterizationFormTests(SimpleTestCase):
    def valid_data(self):
        return {
            "inspection_address": "Rua X", "inspection_date": "2026-09-24",
            "inspector_name": "Responsável",
        }

    def test_new_fields_are_optional(self):
        form = SignalingReportForm(self.valid_data())
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["pgts"], [])
        for name in ("functional_classification", "geometric_configuration", "regulated_speed"):
            self.assertEqual(form.cleaned_data[name], "")

    def test_pgt_collection_type_size_count_and_empty_entries(self):
        field = PGTListField(required=False)
        self.assertEqual(field.clean([" ", " Shopping X ", "", "Escola Y"]), ["Shopping X", "Escola Y"])
        for invalid in ("text", {}, [1], [None], [["nested"]], ["x" * 301], ["x"] * 51):
            with self.subTest(invalid=repr(invalid)[:80]):
                with self.assertRaises(ValidationError):
                    field.clean(invalid)
        self.assertEqual(len(field.clean(["x"] * 50)), 50)

    def test_text_sizes_and_xml_invalid_characters(self):
        for field, maximum in (
            ("functional_classification", 500),
            ("geometric_configuration", 500),
            ("regulated_speed", 100),
        ):
            for value in ("x" * (maximum + 1), "A\x00B", "\uffff", "\ud800"):
                with self.subTest(field=field, value=repr(value)[:40]):
                    form = SignalingReportForm({**self.valid_data(), field: value})
                    self.assertFalse(form.is_valid())
                    self.assertIn(field, form.errors)
        with self.assertRaises(ValidationError):
            PGTListField().clean(["Escola\x0bX"])


class CharacterizationReportTests(TestCase):
    def setUp(self):
        self.point = SignalingPoint.objects.create(
            latitude=-21.17, longitude=-47.81, status="OK", search_radius_meters=100,
        )
        self.url = reverse("signaling:point-report", args=[self.point.pk])
        session = self.client.session
        session[ANALYSIS_SESSION_KEY] = {"view_mode": "individual", "map_data": [], "analysis_id": "current"}
        session.save()
        self.data = {
            "inspection_address": "Rua X", "inspection_date": "2026-09-24",
            "inspector_name": "Responsável",
            "functional_classification": "Via arterial",
            "geometric_configuration": "Cruzamento em quatro ramos",
            "regulated_speed": "50 km/h",
        }

    def docx_text(self, response):
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.streaming)
        content = BytesIO(b"".join(response.streaming_content))
        response.close()
        self.assertTrue(is_zipfile(content))
        with ZipFile(content) as archive:
            for name in archive.namelist():
                if name.endswith(".xml") or name.endswith(".rels"):
                    ElementTree.fromstring(archive.read(name))
        document = Document(content)
        text = "\n".join(
            [p.text for p in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        return text

    @patch("signaling.overpass.urlopen", side_effect=AssertionError("Overpass must not be called"))
    def test_html_card_order_and_opening_never_queries_overpass(self, http):
        response = self.client.get(self.url)
        self.assertContains(response, "Caracterização do local")
        html = response.content.decode()
        self.assertLess(html.index("Informações do estudo"), html.index("Caracterização do local"))
        self.assertLess(html.index("Caracterização do local"), html.index("Filtros aplicados"))
        self.assertContains(response, 'name="functional_classification"')
        self.assertContains(response, 'name="geometric_configuration"')
        self.assertContains(response, 'name="regulated_speed"')
        self.assertContains(response, 'name="pgts"')
        http.assert_not_called()

    @patch("signaling.overpass.urlopen", side_effect=AssertionError("Overpass must not be called"))
    def test_word_uses_final_edited_list_with_unicode_and_literal_xml_text(self, http):
        # A lista enviada já contém a edição e não contém o item removido.
        final_pgts = ["Shopping Center Exemplo", "🏥 Unidade de Saúde: São José & Filhos",
                      "🎓 Unidade escolar: Nome revisado", '<script>alert("X")</script>']
        response = self.client.post(self.url, {**self.data, "pgts": final_pgts})
        text = self.docx_text(response)
        for value in ("Caracterização do local", *self.data.values(), *final_pgts):
            if value != "2026-09-24":
                self.assertIn(value, text)
        self.assertNotIn("Escola removida", text)
        self.assertNotIn("Nome original", text)
        self.assertNotIn("undefined", text)
        self.assertNotIn("None", text)
        http.assert_not_called()

    def test_one_manual_pgt_and_empty_characterization_generate_valid_word(self):
        data = {name: value for name, value in self.data.items()
                if name in ("inspection_address", "inspection_date", "inspector_name")}
        for pgts in (["Shopping X"], [], ["", "  "]):
            with self.subTest(pgts=pgts):
                text = self.docx_text(self.client.post(self.url, {**data, "pgts": pgts}))
                self.assertIn("Caracterização do local", text)
                self.assertIn("Não informado", text)
                self.assertNotIn("None", text)
                self.assertNotIn("undefined", text)
                if pgts == ["Shopping X"]:
                    self.assertIn("Shopping X", text)

    def test_invalid_form_preserves_manual_fields_and_list_with_html_escaping(self):
        pgts = ['</textarea><script>alert("X")</script>', "🎓 Escola editada"]
        data = {**self.data, "inspection_date": "", "pgts": pgts,
                "functional_classification": '<img src=x onerror="alert(1)">'}
        response = self.client.post(self.url, data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context["form"].is_valid())
        self.assertContains(response, "&lt;script&gt;")
        self.assertNotContains(response, '<script>alert("X")</script>')
        self.assertNotContains(response, '<img src=x onerror=')
        self.assertEqual(response.context["form"]["pgts"].value(), pgts)
        self.assertEqual(response.context["form"]["regulated_speed"].value(), "50 km/h")

    def test_invalid_pgt_list_is_retained_and_word_is_not_generated(self):
        values = ["X" * 301, "Manual preservado"]
        with patch("signaling.views.generate_signaling_report") as generate:
            response = self.client.post(self.url, {**self.data, "pgts": values})
        self.assertContains(response, "300 caracteres")
        self.assertContains(response, "Manual preservado")
        self.assertEqual(response.context["form"]["pgts"].value(), values)
        generate.assert_not_called()

    def test_generating_word_does_not_persist_characterization(self):
        before_points = list(SignalingPoint.objects.values())
        before_interventions = list(SignalingIntervention.objects.values())
        before_session = dict(self.client.session)
        self.docx_text(self.client.post(self.url, {**self.data, "pgts": ["Manual"]}))
        self.assertEqual(list(SignalingPoint.objects.values()), before_points)
        self.assertEqual(list(SignalingIntervention.objects.values()), before_interventions)
        self.assertEqual(dict(self.client.session), before_session)
        fresh_page = self.client.get(self.url)
        self.assertEqual(fresh_page.context["form"]["functional_classification"].value(), None)
        self.assertIsNone(fresh_page.context["form"]["pgts"].value())
