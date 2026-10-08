"""Integrações HTTP com CSVs sintéticos e processamento/relatório reais.

Somente o HTTP externo do Overpass é mockado. O polígono municipal versionado,
o banco de teste, a sessão, os templates e o gerador DOCX são usados de verdade.
"""

import csv
import json
import re
from io import BytesIO, StringIO
from unittest.mock import patch
from urllib.parse import parse_qs, urlencode
from xml.etree import ElementTree
from zipfile import ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from docx import Document

from signaling.models import SignalingIntervention, SignalingPoint
from signaling.surveys import ANALYSIS_SESSION_KEY


LATITUDE = "-21.177500"
LONGITUDE = "-47.810300"
SOURCE_FIELDS = (
    "id_sinistro", "data_sinistro", "ano_sinistro", "mes_sinistro",
    "latitude", "longitude", "tp_sinistro_primario", "logradouro",
    "numero_logradouro", "municipio", "tipo_registro", "tipo_via",
    "qtd_gravidade_fatal", "qtd_gravidade_grave", "qtd_gravidade_leve",
    "qtd_gravidade_nao_disponivel", "qtd_gravidade_ileso", "qtd_automovel",
)
ACCEPTED_IDS = {
    "SELECT-JAN", "SELECT-OCT", "OUTSIDE-RADIUS", "WRONG-YEAR",
    "WRONG-MONTH", "WRONG-CATEGORY", "NON-FATAL", "ZERO-SEVERITY",
}


def accident(accident_id, *, year=2022, month=1, **changes):
    row = dict.fromkeys(SOURCE_FIELDS, "")
    row.update({
        "id_sinistro": accident_id,
        "data_sinistro": f"15/{month:02d}/{year}",
        "ano_sinistro": year,
        "mes_sinistro": month,
        "latitude": LATITUDE,
        "longitude": LONGITUDE,
        "tp_sinistro_primario": "COLISAO",
        "logradouro": f"RUA {accident_id}",
        "numero_logradouro": 10,
        "municipio": "RIBEIRAO PRETO",
        "tipo_registro": "SINISTRO FATAL",
        "tipo_via": "VIAS URBANAS",
        "qtd_gravidade_fatal": 1,
        "qtd_automovel": 2,
    })
    row.update(changes)
    return row


def csv_upload(name, rows):
    content = StringIO()
    writer = csv.DictWriter(content, fieldnames=SOURCE_FIELDS, delimiter=";")
    writer.writeheader()
    writer.writerows(rows)
    return SimpleUploadedFile(
        name, content.getvalue().encode("utf-8"), content_type="text/csv",
    )


def individual_files():
    first = [
        accident("SELECT-JAN", tp_sinistro_primario=" Colisão "),
        accident("SELECT-OCT", year=2024, month=10, latitude="-21.177320",
                 tp_sinistro_primario="ATROPELAMENTO"),
        # Aproximadamente 80 m do waypoint; os demais critérios coincidem.
        accident("OUTSIDE-RADIUS", latitude="-21.176780"),
        accident("WRONG-YEAR", year=2025),
        accident("WRONG-MONTH", year=2024, month=2),
        accident("WRONG-CATEGORY", tp_sinistro_primario="CHOQUE"),
        accident("NON-FATAL", tipo_registro="SINISTRO NAO FATAL",
                 qtd_gravidade_fatal="", qtd_gravidade_leve=1),
        accident("ZERO-SEVERITY", tipo_registro="SINISTRO NAO FATAL",
                 qtd_gravidade_fatal=0, qtd_gravidade_ileso=1),
        accident("ONLY-UNINJURED", qtd_gravidade_fatal="", qtd_gravidade_ileso=1),
        accident("INVALID-COORDINATE", latitude="invalid"),
        accident("OTHER-CITY", municipio="SERTAOZINHO"),
        accident("OUTSIDE-MUNICIPALITY", latitude=0, longitude=0),
        accident("OTHER-ROAD", year=2029, tipo_via="RODOVIAS"),
        accident("OTHER-RECORD", tipo_registro="OUTRO REGISTRO"),
    ]
    second = [
        # Se a deduplicação deixar de manter a primeira linha, período e
        # elegibilidade dos dois IDs abaixo mudam.
        accident("SELECT-JAN", year=2030),
        accident("ONLY-UNINJURED"),
        accident("ZERO-UNINJURED", qtd_gravidade_fatal="", qtd_gravidade_ileso=0),
    ]
    return [csv_upload("first.csv", first), csv_upload("second.csv", second)]


class MainFlowFixtures:
    def setUp(self):
        super().setUp()
        http_patch = patch(
            "signaling.overpass.urlopen",
            side_effect=AssertionError("Unexpected external HTTP request"),
        )
        self.http = http_patch.start()
        self.addCleanup(http_patch.stop)

    def upload(self, files, mode="individual"):
        response = self.client.post(
            reverse("analysis"), {"files": files, "view_mode": mode}, follow=True,
        )
        self.assertEqual(response.redirect_chain, [(reverse("analysis"), 302)])
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.context["import_error"])
        return response

    def json_script(self, response, script_id):
        match = re.search(
            rf'<script\b[^>]*\bid="{re.escape(script_id)}"[^>]*>(.*?)</script>',
            response.content.decode(), re.DOTALL,
        )
        self.assertIsNotNone(match, f"Missing JSON payload: {script_id}")
        return json.loads(match.group(1))

    def create_point(self):
        response = self.client.post(
            reverse("signaling:create-point"),
            {"latitude": LATITUDE, "longitude": LONGITUDE, "status": "OK"},
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201)
        return response.json()["id"]

    def read_word(self, response):
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertIn("attachment", response["Content-Disposition"])
        try:
            content = BytesIO(b"".join(response.streaming_content))
        finally:
            response.close()
        with ZipFile(content) as archive:
            for name in archive.namelist():
                if name.endswith((".xml", ".rels")):
                    ElementTree.fromstring(archive.read(name))
        document = Document(content)
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables
               for row in table.rows for cell in row.cells]
        )
        return document, text


class UploadFlowTests(MainFlowFixtures, TestCase):
    def test_csv_consolidation_preparation_and_individual_map_use_real_pipeline(self):
        response = self.upload(individual_files())
        payload = self.json_script(response, "map-data")
        state = self.client.session[ANALYSIS_SESSION_KEY]

        self.assertEqual({item["id"] for item in payload}, ACCEPTED_IDS)
        self.assertEqual(len(payload), len(ACCEPTED_IDS))
        self.assertEqual(payload, state["map_data"])
        selected = next(item for item in payload if item["id"] == "SELECT-JAN")
        self.assertEqual(selected["category"], "collision")
        self.assertTrue(selected["is_fatal"])
        self.assertEqual(selected["year"], 2022)
        self.assertEqual(selected["modes"], [{"name": "Automóvel", "quantity": 2}])
        self.assertEqual(self.json_script(response, "view-mode"), "individual")
        self.assertEqual(self.json_script(response, "analysis-id"), state["analysis_id"])
        self.assertEqual(self.json_script(response, "available-periods"), [
            {"year": 2022, "months": [1]},
            {"year": 2024, "months": [2, 10]},
            {"year": 2025, "months": [1]},
        ])
        summary = state["import_summary"]
        self.assertEqual(summary["file_count"], 2)
        self.assertEqual(summary["accident_count"], 13)
        self.assertEqual(summary["exclusively_uninjured_excluded_count"], 2)
        self.assertEqual(summary["valid_coordinate_count"], 10)
        self.assertEqual(summary["internal_filter_count"], 8)
        self.assertEqual(summary["displayed_count"], 8)
        self.http.assert_not_called()

    def test_eligible_upload_excludes_duplicates_and_uninjured_before_clustering(self):
        first = [
            accident("BELOW-1"), accident("BELOW-2"),
            accident("BELOW-UNINJURED", qtd_gravidade_fatal="", qtd_gravidade_ileso=1),
            *[accident(f"ELIGIBLE-{index}", latitude="-21.180000")
              for index in range(3)],
        ]
        response = self.upload([
            csv_upload("first.csv", first),
            csv_upload("second.csv", [accident("BELOW-1")]),
        ], mode="clusters")

        payload = self.json_script(response, "map-data")
        self.assertEqual(self.json_script(response, "view-mode"), "clusters")
        self.assertEqual(len(payload), 1)
        self.assertAlmostEqual(payload[0]["latitude"], -21.180000)
        self.assertTrue(payload[0]["eligible"])
        self.assertTrue(payload[0]["collision_1y_met"])
        self.assertEqual(payload[0]["collisions_1y"], 3)
        self.assertEqual(payload[0]["period_summary"]["total_count"], 3)
        summary = response.context["import_summary"]
        self.assertEqual(summary["accident_count"], 5)
        self.assertEqual(summary["exclusively_uninjured_excluded_count"], 1)
        self.assertEqual(summary["eligible_count"], 1)
        self.assertEqual(payload, self.client.session[ANALYSIS_SESSION_KEY]["map_data"])
        self.http.assert_not_called()


    def test_upload_with_only_excluded_accidents_renders_empty_map_in_both_modes(self):
        for mode in ("individual", "clusters"):
            with self.subTest(mode=mode):
                response = self.upload([csv_upload("only-uninjured.csv", [
                    accident("EXCLUDED", qtd_gravidade_fatal="", qtd_gravidade_ileso=1),
                ])], mode=mode)
                self.assertEqual(self.json_script(response, "map-data"), [])
                self.assertEqual(self.json_script(response, "available-periods"), [])
                self.assertEqual(self.json_script(response, "view-mode"), mode)
                summary = response.context["import_summary"]
                self.assertEqual(summary["accident_count"], 0)
                self.assertEqual(summary["exclusively_uninjured_excluded_count"], 1)
        self.http.assert_not_called()


class ReportFlowTests(MainFlowFixtures, TestCase):
    def setUp(self):
        super().setUp()
        self.upload(individual_files())
        self.point_id = self.create_point()
        self.report_path = reverse("signaling:point-report", args=[self.point_id])
        self.search_path = reverse("signaling:search-point-pgts", args=[self.point_id])
        self.form_data = {
            "inspection_address": "Rua de teste, 10",
            "inspection_date": "2026-09-30",
            "inspector_name": "Responsável pela vistoria",
        }

    def filtered_url(self):
        state = self.client.session[ANALYSIS_SESSION_KEY]
        return self.report_path + "?" + urlencode([
            ("analysis", state["analysis_id"]),
            ("year", 2022), ("year", 2024),
            ("month", 1), ("month", 10),
            ("category", "collision"), ("category", "pedestrian"),
            ("gravity", "fatal"),
        ])

    def test_uploaded_filters_waypoint_radius_and_widget_render_together(self):
        response = self.client.get(self.filtered_url())
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "signaling/report.html")
        self.assertTemplateUsed(response, "signaling/widgets/pgt_list.html")
        self.assertContains(response, "Caracterização do local")
        self.assertContains(response, 'name="pgts"')
        self.assertInHTML(
            '<textarea class="form-control" rows="1" name="pgts" maxlength="300" '
            'aria-label="Polo Gerador de Tráfego"></textarea>',
            response.content.decode(),
        )
        survey = response.context["survey"]
        self.assertEqual([item["id"] for item in survey["accidents"]],
                         ["SELECT-JAN", "SELECT-OCT"])
        self.assertEqual(survey["radius_meters"], 50)
        self.assertEqual(survey["summary"]["fatal"], 2)
        self.assertEqual(survey["summary"]["non_fatal"], 0)
        self.assertEqual(survey["summary"]["by_type"]["Colisão"], 1)
        self.assertEqual(survey["summary"]["by_type"]["Atropelamento"], 1)
        self.assertContains(response, "Anos: 2022 e 2024. Meses: Janeiro e Outubro.")
        self.assertContains(response, "Somente fatais")

        update = self.client.post(
            reverse("signaling:update-point-search-radius", args=[self.point_id]),
            {"search_radius_meters": 100}, content_type="application/json",
        )
        self.assertEqual(update.status_code, 200)
        wider = self.client.get(self.filtered_url())
        self.assertEqual({item["id"] for item in wider.context["survey"]["accidents"]},
                         {"SELECT-JAN", "SELECT-OCT", "OUTSIDE-RADIUS"})
        self.assertEqual(wider.context["survey"]["radius_meters"], 100)
        map_response = self.client.get(reverse("analysis"))
        point = self.json_script(map_response, "signaling-map-data")[0]
        self.assertEqual((point["id"], point["search_radius_meters"]), (self.point_id, 100))
        self.http.assert_not_called()

    def test_real_word_preserves_filtered_survey_and_transient_characterization(self):
        html = self.client.get(self.filtered_url())
        expected_ids = [item["id"] for item in html.context["survey"]["accidents"]]
        points_before = list(SignalingPoint.objects.values())
        interventions_before = list(SignalingIntervention.objects.values())
        session_before = dict(self.client.session)
        final_pgts = ["Shopping manual", "🎓 Unidade escolar: Nome revisado & <anexo>"]
        characterization = {
            "functional_classification": "Via arterial",
            "geometric_configuration": "Cruzamento em quatro ramos",
            "regulated_speed": "50 km/h",
        }
        document, text = self.read_word(self.client.post(self.filtered_url(), {
            **self.form_data, **characterization, "pgts": final_pgts,
        }))

        accidents_table = next(
            table for table in document.tables if table.cell(0, 0).text == "ID"
        )
        self.assertEqual([row.cells[0].text for row in accidents_table.rows[1:]], expected_ids)
        for value in characterization.values():
            self.assertIn(value, text)
        pgt_values = [
            row.cells[1].text for table in document.tables for row in table.rows
            if row.cells[0].text == "Polos Geradores de Tráfego (PGT)"
        ]
        self.assertEqual(pgt_values, ["\n".join(final_pgts)])
        self.assertIn("Caracterização do local", text)
        for value in html.context["survey"]["filters"].values():
            self.assertIn(value, text)
        self.assertIn("50 m", text)
        self.assertEqual(list(SignalingPoint.objects.values()), points_before)
        self.assertEqual(list(SignalingIntervention.objects.values()), interventions_before)
        self.assertEqual(dict(self.client.session), session_before)
        reloaded = self.client.get(self.filtered_url())
        for name in (*characterization, "pgts"):
            self.assertIsNone(reloaded.context["form"][name].value())
        self.http.assert_not_called()

    def test_overpass_timeout_does_not_block_html_or_manual_word(self):
        session_before = dict(self.client.session)
        self.http.side_effect = TimeoutError("simulated timeout")
        failure = self.client.post(self.search_path)
        self.assertEqual(failure.status_code, 503)
        self.assertFalse(failure.json()["success"])
        self.assertIn("manualmente", failure.json()["error"])
        request = self.http.call_args.args[0]
        query = parse_qs(request.data.decode())["data"][0]
        self.assertIn("around:50,-21.177500,-47.810300", query)

        response = self.client.get(self.filtered_url())
        self.assertContains(response, "Caracterização do local", status_code=200)
        _, text = self.read_word(self.client.post(self.filtered_url(), {
            **self.form_data, "pgts": ["Local manual após falha"],
        }))
        self.assertIn("Local manual após falha", text)
        self.assertIn("SELECT-JAN", text)
        self.assertEqual(self.http.call_count, 1)
        self.assertEqual(dict(self.client.session), session_before)

    def test_empty_overpass_results_and_empty_bound_widget_allow_report(self):
        self.http.side_effect = None
        external_response = self.http.return_value.__enter__.return_value
        external_response.status = 200
        external_response.read.return_value = b'{"elements": []}'
        search = self.client.post(self.search_path)
        self.assertEqual(search.status_code, 200)
        self.assertEqual(search.json(), {"success": True, "results": []})

        html = self.client.get(self.filtered_url())
        self.assertContains(html, "Caracterização do local", status_code=200)
        # Um erro em outro campo deve renderizar o widget com coleção vazia.
        invalid = self.client.post(self.filtered_url(), {"pgts": []})
        self.assertEqual(invalid.status_code, 200)
        self.assertTemplateUsed(invalid, "signaling/widgets/pgt_list.html")
        self.assertEqual(invalid.context["form"]["pgts"].value(), [])
        self.assertContains(invalid, "data-pgt-rows")
        self.assertContains(invalid, "Este campo é obrigatório.")
        _, text = self.read_word(self.client.post(self.filtered_url(), self.form_data))
        self.assertIn("Caracterização do local", text)
        self.assertIn("Não informado", text)
        self.assertEqual(self.http.call_count, 1)

    def test_new_upload_replaces_report_data_and_rejects_old_filter_snapshot(self):
        old_url = self.filtered_url()
        response = self.upload([csv_upload("replacement.csv", [
            accident("REPLACEMENT", year=2026, month=6),
        ])])
        self.assertEqual([item["id"] for item in self.json_script(response, "map-data")],
                         ["REPLACEMENT"])
        self.assertEqual(self.client.get(old_url).status_code, 400)
        self.assertEqual(self.client.post(old_url, self.form_data).status_code, 400)
        current = self.client.get(self.report_path)
        self.assertEqual(current.status_code, 200)
        self.assertEqual([item["id"] for item in current.context["survey"]["accidents"]],
                         ["REPLACEMENT"])
        self.assertTemplateUsed(current, "signaling/widgets/pgt_list.html")
        self.assertEqual(SignalingPoint.objects.count(), 1)
        self.http.assert_not_called()
